import io
import json
import logging
import math
import os
import re
import zipfile
from pathlib import PurePath

from AI_service import ai_service, settings


logger = logging.getLogger(__name__)
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def parse_document(filename: str, data: bytes) -> list[tuple[int | None, str]]:
    extension = PurePath(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Only PDF, DOCX, and TXT files are supported.")
    if not data:
        raise ValueError("The uploaded document is empty.")

    if extension == ".txt":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("TXT files must use UTF-8 encoding.") from exc
        return [(None, text)]

    if extension == ".pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data), strict=True)
            return [
                (page_number, page.extract_text() or "")
                for page_number, page in enumerate(reader.pages, start=1)
            ]
        except Exception as exc:
            raise ValueError("The PDF could not be read.") from exc

    try:
        from docx import Document

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if "word/document.xml" not in archive.namelist():
                raise ValueError("The DOCX file is invalid.")
        document = Document(io.BytesIO(data))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(
                " | ".join(cell.text for cell in row.cells)
                for row in table.rows
            )
        return [(None, "\n".join(paragraphs))]
    except (ValueError, zipfile.BadZipFile):
        raise
    except Exception as exc:
        raise ValueError("The DOCX file could not be read.") from exc


def chunk_pages(
    pages: list[tuple[int | None, str]],
    max_chars: int = settings.MAX_CHUNK_CHARS,
    overlap: int = settings.CHUNK_OVERLAP_CHARS
) -> list[tuple[int | None, str]]:
    if max_chars < 1 or overlap < 0 or overlap >= max_chars:
        raise ValueError("Chunk size and overlap configuration is invalid.")

    chunks: list[tuple[int | None, str]] = []
    for page_number, text in pages:
        normalized = re.sub(r"\s+", " ", text).strip()
        start = 0
        while start < len(normalized):
            end = min(start + max_chars, len(normalized))
            if end < len(normalized):
                boundary = normalized.rfind(" ", start + max_chars // 2, end)
                if boundary > start:
                    end = boundary
            chunk = normalized[start:end].strip()
            if chunk:
                chunks.append((page_number, chunk))
                if len(chunks) > settings.MAX_UPLOAD_CHUNKS:
                    raise ValueError("The document is too long to process.")
            if end >= len(normalized):
                break
            start = max(start + 1, end - overlap)
    return chunks


def create_embedding(text: str, task_type: str) -> list[float]:
    if not os.getenv("GEMINI_API_KEY"):
        raise RuntimeError("GEMINI_API_KEY is not configured.")
    result = ai_service.client.models.embed_content(
        model=settings.GEMINI_EMBEDDING_MODEL,
        contents=text,
        config={"task_type": task_type}
    )
    if not result.embeddings or not result.embeddings[0].values:
        raise RuntimeError("The embedding provider returned no vector.")
    return [float(value) for value in result.embeddings[0].values]


def serialize_embedding(embedding: list[float]) -> str:
    return json.dumps(embedding, separators=(",", ":"), allow_nan=False)


def cosine_similarity(left: list[float], right: list[float]) -> float:
    if len(left) != len(right) or not left:
        return 0.0
    dot_product = sum(a * b for a, b in zip(left, right))
    left_norm = math.sqrt(sum(value * value for value in left))
    right_norm = math.sqrt(sum(value * value for value in right))
    denominator = left_norm * right_norm
    return dot_product / denominator if denominator else 0.0


def search_chunks(
    connection,
    user_id: int,
    query: str,
    limit: int = settings.MAX_RETRIEVED_CHUNKS
) -> list[dict]:
    query_embedding = create_embedding(query, "RETRIEVAL_QUERY")
    rows = connection.execute(
        """
        SELECT chunks.id, chunks.document_id, chunks.page_number,
               chunks.content, chunks.embedding_json, documents.filename
        FROM chunks
        JOIN documents ON documents.id = chunks.document_id
        WHERE documents.user_id = ?
        """,
        (user_id,)
    ).fetchall()
    scored = []
    for row in rows:
        try:
            stored_embedding = json.loads(row["embedding_json"])
            if (
                not isinstance(stored_embedding, list)
                or any(
                    isinstance(value, bool)
                    or not isinstance(value, (int, float))
                    for value in stored_embedding
                )
            ):
                raise ValueError("Stored embedding is not a numeric vector.")
            stored_embedding = [float(value) for value in stored_embedding]
            if not all(math.isfinite(value) for value in stored_embedding):
                raise ValueError("Stored embedding contains a non-finite value.")
        except (json.JSONDecodeError, TypeError, ValueError, OverflowError):
            logger.warning(
                "Skipping chunk with malformed embedding (chunk_id=%s)",
                row["id"]
            )
            continue

        score = cosine_similarity(query_embedding, stored_embedding)
        if score >= settings.MIN_RELEVANCE_SCORE:
            scored.append({
                "chunk_id": row["id"],
                "document_id": row["document_id"],
                "filename": row["filename"],
                "page_number": row["page_number"],
                "content": row["content"],
                "score": score
            })
    scored.sort(key=lambda item: item["score"], reverse=True)
    return scored[:limit]
