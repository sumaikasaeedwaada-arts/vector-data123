import io
import json
import logging
import math
import re
import zipfile
from pathlib import PurePath

import numpy as np

from AI_service import ai_service, settings


logger = logging.getLogger(__name__)
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def parse_document(filename: str, data: bytes) -> list[tuple[int | None, str]]:
    extension = PurePath(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Only PDF, DOCX, and TXT files are supported.")
    if not data:
        raise ValueError("The uploaded document is empty.")

    if extension == ".txt":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("TXT files must use UTF-8 encoding.") from exc
        return [(None, text)]

    if extension == ".pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data), strict=True)
            return [
                (page_number, page.extract_text() or "")
                for page_number, page in enumerate(reader.pages, start=1)
            ]
        except Exception as exc:
            raise ValueError("The PDF could not be read.") from exc

    try:
        from docx import Document

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if "word/document.xml" not in archive.namelist():
                raise ValueError("The DOCX file is invalid.")
        document = Document(io.BytesIO(data))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(
                " | ".join(cell.text for cell in row.cells)
                for row in table.rows
            )
        return [(None, "\n".join(paragraphs))]
    except (ValueError, zipfile.BadZipFile):
        raise
    except Exception as exc:
        raise ValueError("The DOCX file could not be read.") from exc


def chunk_pages(
    pages: list[tuple[int | None, str]],
    max_chars: int = settings.MAX_CHUNK_CHARS,
    overlap: int = settings.CHUNK_OVERLAP_CHARS
) -> list[tuple[int | None, str]]:
    if max_chars < 1 or overlap < 0 or overlap >= max_chars:
        raise ValueError("Chunk size and overlap configuration is invalid.")

    chunks: list[tuple[int | None, str]] = []
    for page_number, text in pages:
        normalized = re.sub(r"\s+", " ", text).strip()
        start = 0
        while start < len(normalized):
            end = min(start + max_chars, len(normalized))
            if end < len(normalized):
                boundary = normalized.rfind(" ", start + max_chars // 2, end)
                if boundary > start:
                    end = boundary
            chunk = normalized[start:end].strip()
            if chunk:
                chunks.append((page_number, chunk))
                if len(chunks) > settings.MAX_UPLOAD_CHUNKS:
                    raise ValueError("The document is too long to process.")
            if end >= len(normalized):
                break
            start = max(start + 1, end - overlap)
    return chunks


def create_embedding(text: str, task_type: str) -> list[float]:
    if not isinstance(text, str) or not text.strip():
        raise ValueError("Embedding text must not be empty.")
    return ai_service.create_embedding(text, task_type)


def validate_embedding(embedding) -> list[float]:
    if not isinstance(embedding, (list, tuple)) or not embedding:
        raise ValueError("Embedding must be a non-empty numeric vector.")
    if any(
        isinstance(value, bool) or not isinstance(value, (int, float))
        for value in embedding
    ):
        raise ValueError("Embedding must contain only numeric values.")
    try:
        vector = [float(value) for value in embedding]
    except (TypeError, ValueError, OverflowError) as exc:
        raise ValueError("Embedding must contain only numeric values.") from exc
    if not all(math.isfinite(value) for value in vector):
        raise ValueError("Embedding must contain only finite numeric values.")
    if not any(vector):
        raise ValueError("Embedding must not be a zero vector.")
    return vector


def serialize_embedding(embedding: list[float]) -> str:
    return json.dumps(
        validate_embedding(embedding),
        separators=(",", ":"),
        allow_nan=False
    )


def initialize_vector_index(connection) -> dict:
    """Verify the persistent SQLite index configuration without clearing data."""
    state = connection.execute(
        "SELECT embedding_model, embedding_dimensions, similarity_metric "
        "FROM vector_index_state WHERE id = 1"
    ).fetchone()
    if state is None:
        raise RuntimeError("The vector index has not been initialized.")
    if state["embedding_model"] != settings.GEMINI_EMBEDDING_MODEL:
        raise RuntimeError(
            "The stored vector index uses a different embedding model; "
            "re-embed its records before searching."
        )
    return {
        "embedding_model": state["embedding_model"],
        "embedding_dimensions": state["embedding_dimensions"],
        "similarity_metric": state["similarity_metric"]
    }


def _ensure_index_dimensions(connection, dimensions: int) -> None:
    state = initialize_vector_index(connection)
    configured_dimensions = state["embedding_dimensions"]
    if configured_dimensions is None:
        connection.execute(
            "UPDATE vector_index_state SET embedding_dimensions = ? WHERE id = 1",
            (dimensions,)
        )
    elif configured_dimensions != dimensions:
        raise ValueError(
            "Embedding dimensions do not match the configured vector index."
        )


def add_embedding(
    connection,
    document_id: str,
    chunk_index: int,
    content: str,
    embedding: list[float],
    page_number: int | None = None
) -> str:
    """Insert or replace a chunk using document ID and position as its stable ID."""
    stored_content = content.strip() if isinstance(content, str) else ""
    if not stored_content:
        raise ValueError("Embedding text must not be empty.")
    if (
        not isinstance(document_id, str)
        or not document_id
        or isinstance(chunk_index, bool)
        or not isinstance(chunk_index, int)
        or chunk_index < 0
    ):
        raise ValueError("A document ID and non-negative chunk index are required.")
    vector = validate_embedding(embedding)
    _ensure_index_dimensions(connection, len(vector))
    connection.execute(
        """
        INSERT INTO chunks
            (document_id, chunk_index, page_number, content, embedding_json)
        VALUES (?, ?, ?, ?, ?)
        ON CONFLICT(document_id, chunk_index) DO UPDATE SET
            page_number = excluded.page_number,
            content = excluded.content,
            embedding_json = excluded.embedding_json
        """,
        (
            document_id,
            chunk_index,
            page_number,
            stored_content,
            serialize_embedding(vector)
        )
    )
    return f"{document_id}:{chunk_index}"


def store_embeddings(connection, document_id: str, chunks: list[tuple]) -> int:
    """Atomically upsert chunk content, metadata, and vectors for one document."""
    if not isinstance(document_id, str) or not document_id:
        raise ValueError("A document ID is required.")
    prepared = []
    for chunk_index, (page_number, content, embedding) in enumerate(chunks):
        stored_content = content.strip() if isinstance(content, str) else ""
        if not stored_content:
            raise ValueError("Embedding text must not be empty.")
        vector = validate_embedding(embedding)
        prepared.append((chunk_index, page_number, stored_content, vector))
    if not prepared:
        return 0

    dimensions = len(prepared[0][3])
    if any(len(item[3]) != dimensions for item in prepared):
        raise ValueError("Embedding dimensions are inconsistent.")
    _ensure_index_dimensions(connection, dimensions)
    connection.executemany(
        """
        INSERT INTO chunks
            (document_id, chunk_index, page_number, content, embedding_json)
        VALUES (?, ?, ?, ?, ?)
        ON CONFLICT(document_id, chunk_index) DO UPDATE SET
            page_number = excluded.page_number,
            content = excluded.content,
            embedding_json = excluded.embedding_json
        """,
        [
            (
                document_id,
                chunk_index,
                page_number,
                content,
                serialize_embedding(embedding)
            )
            for chunk_index, page_number, content, embedding in prepared
        ]
    )
    return len(prepared)


def get_embedding(
    connection,
    document_id: str,
    chunk_index: int,
    user_id: int | None = None
) -> dict | None:
    query = """
        SELECT chunks.id, chunks.document_id, chunks.chunk_index,
               chunks.page_number, chunks.content, chunks.embedding_json,
               documents.filename, documents.user_id
        FROM chunks
        JOIN documents ON documents.id = chunks.document_id
        WHERE chunks.document_id = ? AND chunks.chunk_index = ?
    """
    parameters: list[object] = [document_id, chunk_index]
    if user_id is not None:
        query += " AND documents.user_id = ?"
        parameters.append(user_id)
    row = connection.execute(query, parameters).fetchone()
    if row is None:
        return None
    embedding = json.loads(row["embedding_json"])
    return {
        "id": f'{row["document_id"]}:{row["chunk_index"]}',
        "chunk_id": row["id"],
        "document_id": row["document_id"],
        "chunk_index": row["chunk_index"],
        "page_number": row["page_number"],
        "content": row["content"],
        "embedding": validate_embedding(embedding),
        "metadata": {
            "document_id": row["document_id"],
            "source": row["filename"],
            "filename": row["filename"],
            "user_id": row["user_id"],
            "page_number": row["page_number"]
        }
    }


def update_embedding(
    connection,
    document_id: str,
    chunk_index: int,
    content: str,
    embedding: list[float],
    page_number: int | None = None
) -> str:
    return add_embedding(
        connection,
        document_id,
        chunk_index,
        content,
        embedding,
        page_number
    )


def delete_embedding(
    connection,
    document_id: str,
    chunk_index: int,
    user_id: int | None = None
) -> bool:
    query = """
        DELETE FROM chunks
        WHERE document_id = ? AND chunk_index = ?
        AND document_id IN (
            SELECT id FROM documents
            WHERE id = ? {user_filter}
        )
    """
    user_filter = "AND user_id = ?" if user_id is not None else ""
    query = query.format(user_filter=user_filter)
    parameters: list[object] = [document_id, chunk_index, document_id]
    if user_id is not None:
        parameters.append(user_id)
    return connection.execute(query, parameters).rowcount > 0


def get_index_status(connection, user_id: int | None = None) -> dict:
    status = initialize_vector_index(connection)
    if user_id is None:
        row = connection.execute("SELECT COUNT(*) AS record_count FROM chunks").fetchone()
    else:
        row = connection.execute(
            """
            SELECT COUNT(*) AS record_count
            FROM chunks
            JOIN documents ON documents.id = chunks.document_id
            WHERE documents.user_id = ?
            """,
            (user_id,)
        ).fetchone()
    return {**status, "record_count": row["record_count"]}


def cosine_similarity(left: list[float], right: list[float]) -> float:
    if len(left) != len(right) or not left:
        return 0.0
    dot_product = sum(a * b for a, b in zip(left, right))
    left_norm = math.sqrt(sum(value * value for value in left))
    right_norm = math.sqrt(sum(value * value for value in right))
    denominator = left_norm * right_norm
    return dot_product / denominator if denominator else 0.0


def search_chunks(
    connection,
    user_id: int,
    query: str,
    limit: int = settings.MAX_RETRIEVED_CHUNKS,
    filters: dict | None = None
) -> list[dict]:
    if not isinstance(query, str) or not query.strip():
        raise ValueError("Search query must not be empty.")
    if isinstance(limit, bool) or not isinstance(limit, int) or limit < 1:
        raise ValueError("Search limit must be a positive integer.")

    filter_columns = {
        "document_id": "documents.id",
        "filename": "documents.filename",
        "page_number": "chunks.page_number"
    }
    where = ["documents.user_id = ?"]
    parameters: list[object] = [user_id]
    if filters:
        for key, value in filters.items():
            if key not in filter_columns:
                raise ValueError(f"Unsupported vector metadata filter: {key}")
            if value is None:
                where.append(f"{filter_columns[key]} IS NULL")
            else:
                where.append(f"{filter_columns[key]} = ?")
                parameters.append(value)

    query_rows = connection.execute(
        f"""
        SELECT chunks.id, chunks.document_id, chunks.chunk_index,
               chunks.page_number, chunks.content, chunks.embedding_json,
               documents.filename, documents.user_id
        FROM chunks
        JOIN documents ON documents.id = chunks.document_id
        WHERE {" AND ".join(where)}
        """,
        parameters
    ).fetchall()
    if not query_rows:
        return []

    query_embedding = create_embedding(query, "RETRIEVAL_QUERY")
    query_embedding = validate_embedding(query_embedding)
    index_state = initialize_vector_index(connection)
    configured_dimensions = index_state["embedding_dimensions"]
    if (
        configured_dimensions is not None
        and len(query_embedding) != configured_dimensions
    ):
        raise ValueError(
            "Query embedding dimensions do not match the configured vector index."
        )
    valid_rows = []
    vectors = []
    for row in query_rows:
        try:
            stored_embedding = json.loads(row["embedding_json"])
            stored_embedding = validate_embedding(stored_embedding)
        except (json.JSONDecodeError, TypeError, ValueError, OverflowError):
            logger.warning(
                "Skipping chunk with malformed embedding (chunk_id=%s)",
                row["id"]
            )
            continue

        if len(stored_embedding) != len(query_embedding):
            logger.warning(
                "Skipping chunk with mismatched embedding dimensions (chunk_id=%s)",
                row["id"]
            )
            continue
        valid_rows.append(row)
        vectors.append(stored_embedding)

    if not vectors:
        return []
    matrix = np.asarray(vectors, dtype=np.float64)
    query_vector = np.asarray(query_embedding, dtype=np.float64)
    denominators = np.linalg.norm(matrix, axis=1) * np.linalg.norm(query_vector)
    scores = np.divide(
        matrix @ query_vector,
        denominators,
        out=np.zeros(len(vectors), dtype=np.float64),
        where=denominators != 0
    )
    selected = np.flatnonzero(scores >= settings.MIN_RELEVANCE_SCORE)
    if selected.size > limit:
        selected = selected[np.argpartition(scores[selected], -limit)[-limit:]]
    selected = selected[np.argsort(scores[selected])[::-1]]

    results = []
    for index in selected:
        row = valid_rows[int(index)]
        results.append({
            "id": f'{row["document_id"]}:{row["chunk_index"]}',
            "chunk_id": row["id"],
            "document_id": row["document_id"],
            "filename": row["filename"],
            "page_number": row["page_number"],
            "content": row["content"],
            "metadata": {
                "document_id": row["document_id"],
                "source": row["filename"],
                "filename": row["filename"],
                "user_id": row["user_id"],
                "page_number": row["page_number"]
            },
            "score": float(scores[int(index)])
        })
    return results
